import os
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events():
    """Seed the events table from data/events_seed.xlsx if not already seeded."""
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    row_count = cursor.fetchone()["count"]
    if row_count > 0:
        conn.close()
        return False, f"Events table already contains {row_count} rows. Seeding skipped."

    # Validate that data/events_seed.xlsx exists
    if not SEED_FILE_PATH.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed file not found at {SEED_FILE_PATH}. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    # Load Excel workbook
    workbook = openpyxl.load_workbook(str(SEED_FILE_PATH), data_only=True)
    if "Events" not in workbook.sheetnames:
        conn.close()
        raise ValueError(
            "The sheet 'Events' was not found in data/events_seed.xlsx. "
            "Ensure the companion seed spreadsheet has the 'Events' sheet."
        )

    sheet = workbook["Events"]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows or len(rows) < 2:
        conn.close()
        raise ValueError("Seed spreadsheet 'Events' sheet has no data rows.")

    # Row 1 is header: id, title, category, description, event_date, event_time, venue, capacity, fee, icon
    header = rows[0]
    expected_headers = [
        "id", "title", "category", "description", "event_date",
        "event_time", "venue", "capacity", "fee", "icon"
    ]

    # Insert rows into events table
    inserted_count = 0
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue
        event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id) if event_id is not None else None,
                str(title or "").strip(),
                str(category or "").strip(),
                str(description or "").strip() if description is not None else None,
                str(event_date or "").strip(),
                str(event_time or "").strip(),
                str(venue or "").strip(),
                int(capacity) if capacity is not None else 0,
                int(fee) if fee is not None else 0,
                str(icon or "").strip() if icon is not None else None,
            ),
        )
        inserted_count += 1

    conn.commit()
    conn.close()
    return True, f"Successfully seeded {inserted_count} events from {SEED_FILE_PATH}."


if __name__ == "__main__":
    status, msg = seed_events()
    print(msg)
    conn = get_db_connection()
    events = conn.execute("SELECT id, title, category, venue, capacity, fee, icon FROM events ORDER BY id").fetchall()
    print(f"\nTotal events in database: {len(events)}")
    for e in events:
        print(f"[{e['id']}] {e['icon']} {e['title']} ({e['category']}) - Venue: {e['venue']}, Capacity: {e['capacity']}, Fee: Rs.{e['fee']}")
    conn.close()
